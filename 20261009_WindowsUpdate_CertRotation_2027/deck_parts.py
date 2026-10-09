"""Shared building blocks for the 2026-10-09 short news decks."""
import sys

sys.path.insert(0, "/home/ebi/.claude/skills/youtube-slide/scripts")

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Pt  # noqa: E402

from slide_helpers import (  # noqa: E402,F401
    BLUE, CARD_BORDER, CARD_FILL, EMU, GREEN, LABEL, ORANGE, PURPLE, RED, TEAL,
    TEXT_MAIN, TEXT_MUTED, add_slide, box, header, new_deck, text, title_slide,
)


def card(slide, x, y, w, h, title, body, accent=BLUE, title_size=22, body_size=20):
    box(slide, x, y, w, h, fill=CARD_FILL, border=CARD_BORDER, bw=1.2)
    box(slide, x, y, 0.08 * EMU, h, fill=accent, border=accent, radius=0.01)
    text(slide, x + 0.25 * EMU, y + 0.15 * EMU, w - 0.45 * EMU, 0.6 * EMU,
         [{"text": title, "size": title_size, "bold": True, "color": accent}])
    text(slide, x + 0.25 * EMU, y + 0.8 * EMU, w - 0.45 * EMU, h - 0.9 * EMU,
         [{"text": ln, "size": body_size, "color": TEXT_MAIN, "space_after": 6}
          for ln in body.split("\n")])


def big_statement(prs, num, label, top, main, sub, color=ORANGE, main_size=64):
    s = add_slide(prs)
    text(s, 0.48 * EMU, 0.38 * EMU, 0.8 * EMU, 0.4 * EMU,
         [{"text": num, "size": 18, "bold": True, "color": color}])
    text(s, 1.15 * EMU, 0.4 * EMU, 6 * EMU, 0.4 * EMU,
         [{"text": label, "size": 13, "color": LABEL}])
    text(s, 0.6 * EMU, 1.4 * EMU, 12.1 * EMU, 0.6 * EMU,
         [{"text": top, "size": 26, "color": TEXT_MUTED}], align=PP_ALIGN.CENTER)
    text(s, 0.4 * EMU, 2.1 * EMU, 12.5 * EMU, 1.9 * EMU,
         [{"text": main, "size": main_size, "bold": True, "color": color}],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.9 * EMU, 4.4 * EMU, 11.5 * EMU, 2.2 * EMU,
         [{"text": ln, "size": 26, "color": TEXT_MAIN, "space_after": 8}
          for ln in sub.split("\n")], align=PP_ALIGN.CENTER)
    return s


def shrink_title(slide, size=40):
    """Shrink the title_slide headline so long product names stay on one line."""
    for shape in slide.shapes:
        if shape.has_text_frame:
            for run in (r for para in shape.text_frame.paragraphs for r in para.runs):
                if run.font.size and run.font.size.pt == 46:
                    run.font.size = Pt(size)


def checklist(prs, num, label, title, items, color=GREEN, step=1.5, height=1.3):
    s = add_slide(prs)
    header(s, num, label, title)
    y = 2.0
    for i, item in enumerate(items, 1):
        box(s, 0.7 * EMU, y * EMU, 11.9 * EMU, height * EMU, fill=CARD_FILL, border=CARD_BORDER)
        text(s, 0.95 * EMU, y * EMU, 0.9 * EMU, height * EMU,
             [{"text": f"□{i}", "size": 30, "bold": True, "color": color}], anchor=MSO_ANCHOR.MIDDLE)
        text(s, 1.95 * EMU, y * EMU, 10.5 * EMU, height * EMU,
             [{"text": item, "size": 24, "color": TEXT_MAIN}], anchor=MSO_ANCHOR.MIDDLE)
        y += step
    return s
